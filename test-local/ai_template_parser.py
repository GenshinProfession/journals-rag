"""
AI模板解析服务
使用阿里云 qwen-plus 模型解析论文模板，提取格式规范
"""

import json
import requests
from pathlib import Path
from typing import Optional

# 阿里云模型配置
API_KEY = "sk-2aa7503c50c14f96b408d27911b8120d"
MODEL = "qwen-plus-2025-07-28"
API_URL = "https://dashscope.aliyuncs.com/compatible-mode/v1/chat/completions"


def convert_doc_to_docx(doc_path: str) -> Optional[str]:
    """将 .doc 转换为 .docx"""
    import subprocess
    from pathlib import Path
    
    doc_path = Path(doc_path)
    docx_path = doc_path.parent / f"{doc_path.stem}_converted.docx"
    
    # 使用 wps_convert.py 转换
    wps_script = Path(__file__).parent / "wps_convert.py"
    
    try:
        result = subprocess.run(
            ['python', str(wps_script), str(doc_path), str(docx_path)],
            capture_output=True,
            text=True,
            timeout=120
        )
        
        if result.returncode == 0 and docx_path.exists():
            return str(docx_path)
        return None
    except Exception as e:
        print(f"转换失败: {e}")
        return None


def extract_text_from_file(file_path: str) -> str:
    """从 doc 或 docx 文件提取文本"""
    file_path = Path(file_path)
    
    # 如果是 .doc 文件，先转换为 .docx
    if file_path.suffix.lower() == '.doc':
        print("  检测到 .doc 文件，正在转换为 .docx...")
        docx_path = convert_doc_to_docx(str(file_path))
        if docx_path:
            file_path = Path(docx_path)
            print(f"  转换成功: {file_path}")
        else:
            return "转换失败: 无法将 .doc 转换为 .docx"
    
    try:
        from docx import Document
        doc = Document(str(file_path))
        text_parts = []
        
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        
        # 也提取表格中的文本
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)
        
        return "\n".join(text_parts)
    except Exception as e:
        return f"提取文本失败: {e}"


def call_ai_model(prompt: str, system_prompt: str = "") -> str:
    """调用阿里云模型"""
    headers = {
        "Authorization": f"Bearer {API_KEY}",
        "Content-Type": "application/json"
    }
    
    messages = []
    if system_prompt:
        messages.append({"role": "system", "content": system_prompt})
    messages.append({"role": "user", "content": prompt})
    
    data = {
        "model": MODEL,
        "messages": messages,
        "temperature": 0.1,  # 低温度，确保输出稳定
        "max_tokens": 4000
    }
    
    try:
        response = requests.post(API_URL, headers=headers, json=data, timeout=60)
        response.raise_for_status()
        result = response.json()
        return result["choices"][0]["message"]["content"]
    except Exception as e:
        return f"AI调用失败: {e}"


SYSTEM_PROMPT = """你是一个论文格式规范解析专家。你的任务是从论文模板文档中提取格式规范，并生成标准的JSON配置文件。

你需要提取以下信息：
1. 页面设置（纸张大小、方向）
2. 页边距（上、下、左、右、装订线）
3. 封面格式（中英文标题、字段的字体字号）
4. 摘要格式（中英文标题、正文的字体字号行距）
5. 关键词格式（字体、字号、分隔符）
6. 目录格式（标题、条目的字体字号）
7. 正文格式（字体、字号、行距、首行缩进）
8. 标题格式（一级到四级标题的字体字号）
9. 图表格式（标题的字体字号位置）
10. 参考文献格式（字体、字号、行距、引用格式）
11. 附录格式
12. 致谢格式

请严格按照以下JSON结构输出，不要输出其他内容：

```json
{
  "_metadata": {
    "version": "2.0",
    "school_name": "学校名称",
    "degree_level": "bachelor/master/doctor",
    "parsed_at": "解析时间"
  },
  "page": {
    "size": "A4/A3",
    "orientation": "portrait/landscape"
  },
  "margin": {
    "top": 2.54,
    "bottom": 2.54,
    "left": 3.17,
    "right": 3.17,
    "binding": 0
  },
  "cover": {
    "title": {"font": "字体", "size": "字号", "bold": true/false, "align": "居中/左对齐"},
    "field": {"font": "字体", "size": "字号"},
    "date_format": "日期格式"
  },
  "abstract_cn": {
    "title": {"font": "字体", "size": "字号", "bold": true/false, "align": "居中"},
    "body": {"font": "字体", "size": "字号", "line_spacing": "行距", "first_indent": "缩进"},
    "max_words": 0
  },
  "abstract_en": {
    "title": {"font": "字体", "size": "字号", "bold": true/false, "align": "居中"},
    "body": {"font": "字体", "size": "字号", "line_spacing": "行距"}
  },
  "keywords_cn": {"font": "字体", "size": "字号", "separator": "分隔符", "count_min": 3, "count_max": 5},
  "keywords_en": {"font": "字体", "size": "字号", "separator": "分隔符"},
  "toc": {
    "title": {"font": "字体", "size": "字号", "bold": true/false},
    "level1": {"font": "字体", "size": "字号"},
    "level2": {"font": "字体", "size": "字号"},
    "level3": {"font": "字体", "size": "字号"}
  },
  "body": {
    "font": "字体",
    "size": "字号",
    "line_spacing": "行距",
    "first_indent": "缩进",
    "align": "对齐方式"
  },
  "heading1": {"font": "字体", "size": "字号", "bold": true/false, "align": "居中/左对齐", "spacing_before": "段前", "spacing_after": "段后"},
  "heading2": {"font": "字体", "size": "字号", "bold": true/false, "align": "居中/左对齐", "spacing_before": "段前", "spacing_after": "段后"},
  "heading3": {"font": "字体", "size": "字号", "bold": true/false, "align": "居中/左对齐", "spacing_before": "段前", "spacing_after": "段后"},
  "heading4": {"font": "字体", "size": "字号", "bold": true/false, "align": "居中/左对齐", "spacing_before": "段前", "spacing_after": "段后"},
  "figure": {
    "title": {"font": "字体", "size": "字号", "bold": true/false, "align": "居中"},
    "position": "图下方/图上方"
  },
  "table": {
    "title": {"font": "字体", "size": "字号", "bold": true/false, "align": "居中"},
    "position": "表上方/表下方"
  },
  "reference": {
    "style": "GB/T 7714/APA/MLA",
    "font": "字体",
    "size": "字号",
    "line_spacing": "行距",
    "numbering": "[1]/1./(1)"
  },
  "appendix": {
    "title": {"font": "字体", "size": "字号", "bold": true/false},
    "body": {"font": "字体", "size": "字号", "line_spacing": "行距"}
  },
  "acknowledgement": {
    "title": {"font": "字体", "size": "字号", "bold": true/false},
    "body": {"font": "字体", "size": "字号", "line_spacing": "行距"}
  },
  "header": {"font": "字体", "size": "字号", "content": "内容"},
  "footer": {"font": "字体", "size": "字号", "numbering": "阿拉伯数字/罗马数字"}
}
```

注意：
1. 如果模板中没有明确说明某个字段，使用合理的默认值
2. 字号使用中文表示（如：小四号、四号、小二号）
3. 行距使用倍数表示（如：1.5倍、2倍、固定值20磅）
4. 页边距单位是厘米
5. 确保输出的是有效的JSON格式"""


def parse_template(file_path: str) -> dict:
    """解析模板文档，返回格式规范JSON"""
    
    # 1. 提取文本
    print("[1/3] 提取文档文本...")
    text = extract_text_from_file(file_path)
    
    if text.startswith("提取文本失败"):
        return {"error": text}
    
    print(f"  提取到 {len(text)} 字符")
    
    # 2. 调用AI解析
    print("[2/3] 调用AI解析...")
    prompt = f"""请解析以下论文模板文档，提取格式规范并生成JSON配置：

---文档内容开始---
{text[:8000]}
---文档内容结束---

请严格按照系统提示中的JSON格式输出，不要有任何额外说明。"""
    
    result = call_ai_model(prompt, SYSTEM_PROMPT)
    
    if result.startswith("AI调用失败"):
        return {"error": result}
    
    # 3. 解析JSON
    print("[3/3] 解析AI输出...")
    try:
        # 提取JSON部分（可能被包裹在 ```json ``` 中）
        json_str = result
        if "```json" in json_str:
            json_str = json_str.split("```json")[1].split("```")[0]
        elif "```" in json_str:
            json_str = json_str.split("```")[1].split("```")[0]
        
        json_str = json_str.strip()
        spec = json.loads(json_str)
        return spec
    except json.JSONDecodeError as e:
        return {
            "error": f"JSON解析失败: {e}",
            "raw_output": result
        }


def save_spec(spec: dict, output_path: str):
    """保存格式规范到JSON文件"""
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(spec, f, ensure_ascii=False, indent=2)
    print(f"格式规范已保存到: {output_path}")


if __name__ == "__main__":
    import sys
    
    if len(sys.argv) < 2:
        print("用法: python ai_template_parser.py <模板文档路径> [输出JSON路径]")
        print("示例: python ai_template_parser.py 论文模板.docx 格式规范.json")
        sys.exit(1)
    
    input_path = sys.argv[1]
    output_path = sys.argv[2] if len(sys.argv) > 2 else "格式规范.json"
    
    if not Path(input_path).exists():
        print(f"错误: 文件不存在 {input_path}")
        sys.exit(1)
    
    print(f"开始解析模板: {input_path}")
    print("=" * 50)
    
    spec = parse_template(input_path)
    
    if "error" in spec:
        print(f"错误: {spec['error']}")
        if "raw_output" in spec:
            print(f"AI原始输出:\n{spec['raw_output']}")
        sys.exit(1)
    
    save_spec(spec, output_path)
    print("=" * 50)
    print("解析完成!")

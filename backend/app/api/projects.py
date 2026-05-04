import json
from io import BytesIO
from uuid import UUID

from fastapi import APIRouter, HTTPException, Query, Response
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.deps import DbSessionDep, EmbeddingServiceDep, LLMServiceDep, SettingsDep, WriterUserDep
from app.models.ai_usage import AIUsageRecord
from app.models.model_catalog import ModelCatalog
from app.models.project import Chapter, Project
from app.models.rag import Literature, ReferenceReview
from app.models.school import School, SchoolTemplateGroup
from app.models.user import User
from app.schemas.projects import (
    ChapterGenerateRequest,
    ChapterResponse,
    ChapterReviewRequest,
    ChapterRewriteRequest,
    ChapterUpdate,
    OutlineGenerateRequest,
    ProjectCreate,
    ProjectResponse,
)
from app.schemas.rag import ReferenceReviewRequest
from app.services.literature_text import resolve_literature_text
from app.services.vector_store import VectorStore

router = APIRouter()


def _writer_project(db: Session, writer: User, project_id: UUID) -> Project:
    project = db.get(Project, project_id)
    if project is None or project.user_id != writer.id:
        raise HTTPException(status_code=404, detail="Project not found")
    return project


def _enabled_model_for_scenario(
    db: Session, model_id: UUID | None, project: Project, scenario: str
) -> ModelCatalog:
    candidate_id = model_id or project.default_model_id
    model = db.get(ModelCatalog, candidate_id) if candidate_id else None
    if model is None:
        stmt = (
            select(ModelCatalog)
            .where(ModelCatalog.enabled.is_(True))
            .order_by(ModelCatalog.sort_order, ModelCatalog.display_name)
        )
        for row in db.scalars(stmt).all():
            allowed = row.allowed_scenarios or []
            if not allowed or scenario in allowed:
                model = row
                break
    if model is None or not model.enabled:
        raise HTTPException(status_code=400, detail=f"No enabled model available for {scenario}")
    allowed = model.allowed_scenarios or []
    if allowed and scenario not in allowed:
        raise HTTPException(status_code=400, detail=f"Model not allowed for {scenario}")
    return model


def _parse_json_object(raw: str) -> dict[str, object]:
    if not raw.strip():
        return {}
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        start = raw.find("{")
        end = raw.rfind("}")
        if start < 0 or end <= start:
            return {}
        try:
            data = json.loads(raw[start : end + 1])
        except json.JSONDecodeError:
            return {}
    return data if isinstance(data, dict) else {}


def _latest_usage_id(db: Session, user_id: UUID) -> UUID | None:
    stmt = (
        select(AIUsageRecord.id)
        .where(AIUsageRecord.user_id == user_id)
        .order_by(AIUsageRecord.created_at.desc())
        .limit(1)
    )
    return db.scalar(stmt)


def _school_context(db: Session, project: Project) -> str:
    if project.school_id is None:
        return "未选择学校模板。"
    group = db.get(SchoolTemplateGroup, project.school_id)
    if group is None:
        return "学校模板不存在或已删除。"
    school = db.get(School, group.school_id)
    school_name = school.name if school else "未知"
    parts = [
        f"学校：{school_name}",
        f"层次：{group.degree_level}",
        f"专业：{group.discipline or '通用'}",
        f"年份：{group.year or '通用'}",
        f"引用格式：{group.citation_style or '未指定'}",
    ]
    if group.structure and group.structure.structure_json:
        parts.append(f"结构DSL：{json.dumps(group.structure.structure_json, ensure_ascii=False)}")
    if group.format_rules and group.format_rules.rules_json:
        parts.append(f"格式DSL：{json.dumps(group.format_rules.rules_json, ensure_ascii=False)}")
    if group.citation_rules:
        if group.citation_rules.citation_json:
            parts.append(f"引用DSL：{json.dumps(group.citation_rules.citation_json, ensure_ascii=False)}")
        if group.citation_rules.citation_text:
            parts.append(f"引用模板：\n{group.citation_rules.citation_text}")
    return "\n".join(parts)


def _reference_lines(db: Session, project_id: UUID) -> list[str]:
    stmt = select(Literature).where(Literature.project_id == project_id).order_by(Literature.year, Literature.title)
    rows = list(db.scalars(stmt).all())
    lines: list[str] = []
    for item in rows:
        year = f" ({item.year})" if item.year else ""
        authors = f"{item.authors}. " if item.authors else ""
        journal = f" {item.journal}." if item.journal else ""
        doi = f" DOI: {item.doi}." if item.doi else ""
        lines.append(f"{authors}{item.title}{year}.{journal}{doi}".strip())
    return lines


@router.get("", response_model=list[ProjectResponse])
def list_projects(writer: WriterUserDep, db: DbSessionDep) -> list[Project]:
    stmt = (
        select(Project)
        .where(Project.user_id == writer.id)
        .order_by(Project.updated_at.desc())
    )
    return list(db.scalars(stmt).all())


@router.post("", response_model=ProjectResponse, status_code=201)
def create_project(writer: WriterUserDep, db: DbSessionDep, payload: ProjectCreate) -> Project:
    if payload.school_id is not None:
        group = db.get(SchoolTemplateGroup, payload.school_id)
        if group is None or not group.enabled:
            raise HTTPException(status_code=400, detail="School template group not found or disabled")
    project = Project(
        user_id=writer.id,
        school_id=payload.school_id,
        degree_level=payload.degree_level,
        discipline=payload.discipline,
        title=payload.title,
        topic=payload.topic,
        default_model_id=payload.default_model_id,
    )
    db.add(project)
    db.commit()
    db.refresh(project)
    return project


@router.post("/{project_id}/reference/review")
async def review_reference(
    writer: WriterUserDep,
    db: DbSessionDep,
    settings: SettingsDep,
    llm: LLMServiceDep,
    project_id: UUID,
    payload: ReferenceReviewRequest,
) -> dict[str, object]:
    project = _writer_project(db, writer, project_id)
    literature = db.get(Literature, payload.literature_id)
    if literature is None or literature.project_id != project.id:
        raise HTTPException(status_code=404, detail="Literature not found")
    model = _enabled_model_for_scenario(db, payload.model_id, project, "reference_review")

    paper_text = resolve_literature_text(literature, settings.upload_root)[:24000]
    if not paper_text:
        raise HTTPException(status_code=400, detail="reference_has_no_extractable_text")

    prompt = (
        "你是论文项目的标准参考论文质检员。请评估这篇文献是否可作为后续 RAG "
        "与论文写作的标准参考论文。只输出 JSON，字段：topic_relevance_score, "
        "structure_score, academic_quality_score, overall_score, passed, issues, report。"
    )
    content = await llm.call(
        user_id=writer.id,
        project_id=project.id,
        agent_name="reference_review",
        scenario="reference_review",
        model=model,
        messages=[
            {"role": "system", "content": prompt},
            {
                "role": "user",
                "content": (
                    f"项目：{project.degree_level} / {project.discipline}\n"
                    f"题目：{project.title or project.topic or '未定'}\n"
                    f"文献标题：{literature.title}\n\n正文：\n{paper_text}"
                ),
            },
        ],
    )
    data = _parse_json_object(content)
    if not data:
        data = {
            "topic_relevance_score": 70,
            "structure_score": 70,
            "academic_quality_score": 70,
            "overall_score": 70,
            "passed": True,
            "issues": [],
            "report": "开发模式或模型未返回 JSON，已按可继续入库的默认审核结果记录。",
        }

    review = ReferenceReview(
        project_id=project.id,
        user_id=writer.id,
        literature_id=literature.id,
        model_id=model.id,
        topic_relevance_score=int(data.get("topic_relevance_score") or 0),
        structure_score=int(data.get("structure_score") or 0),
        academic_quality_score=int(data.get("academic_quality_score") or 0),
        overall_score=int(data.get("overall_score") or 0),
        passed=bool(data.get("passed")),
        issues=data.get("issues") if isinstance(data.get("issues"), list) else [],
        report=str(data.get("report") or ""),
        usage_record_id=_latest_usage_id(db, writer.id),
    )
    db.add(review)
    literature.rag_status = "review_passed" if review.passed else "review_failed"
    db.commit()
    db.refresh(review)
    return {
        "status": "completed",
        "project_id": str(project.id),
        "review_id": str(review.id),
        "literature_id": str(literature.id),
        "model_id": str(model.id),
        "passed": review.passed,
        "overall_score": review.overall_score,
        "report": review.report,
    }


@router.post("/{project_id}/outline/generate")
async def generate_outline(
    writer: WriterUserDep,
    db: DbSessionDep,
    llm: LLMServiceDep,
    project_id: UUID,
    payload: OutlineGenerateRequest,
) -> dict[str, object]:
    project = _writer_project(db, writer, project_id)
    model = _enabled_model_for_scenario(db, payload.model_id, project, "outline")

    content = await llm.call(
        user_id=writer.id,
        project_id=project.id,
        agent_name="outline_generator",
        scenario="outline",
        model=model,
        messages=[
            {
                "role": "system",
                "content": (
                    "你是论文大纲规划 Agent。只输出 JSON："
                    "{\"chapters\":[{\"title\":\"...\",\"summary\":\"...\"}]}"
                ),
            },
            {
                "role": "user",
                "content": (
                    f"层次：{project.degree_level}\n"
                    f"学科：{project.discipline}\n"
                    f"题目：{project.title or '未定'}\n"
                    f"主题：{project.topic or '未定'}\n\n"
                    f"学校/格式模板：\n{_school_context(db, project)}"
                ),
            },
        ],
    )
    data = _parse_json_object(content)
    chapters = data.get("chapters")
    if not isinstance(chapters, list) or not chapters:
        chapters = [
            {"title": "绪论", "summary": "研究背景、意义、问题与方法。"},
            {"title": "文献综述", "summary": "梳理国内外研究现状与理论基础。"},
            {"title": "研究设计", "summary": "说明研究对象、数据来源与方法。"},
            {"title": "分析与讨论", "summary": "展开核心分析并回应研究问题。"},
            {"title": "结论与建议", "summary": "总结发现、局限与后续展望。"},
        ]

    project.outline = {"chapters": chapters}
    project.status = "outline_ready"

    existing = list(
        db.scalars(select(Chapter).where(Chapter.project_id == project.id).limit(1)).all()
    )
    if not existing:
        for idx, item in enumerate(chapters):
            if not isinstance(item, dict):
                continue
            title = str(item.get("title") or f"第 {idx + 1} 章")
            summary = str(item.get("summary") or "")
            db.add(
                Chapter(
                    project_id=project.id,
                    title=title[:255],
                    order_index=idx,
                    content=summary,
                    word_count=len(summary),
                    status="draft",
                )
            )
    db.commit()
    return {
        "status": "completed",
        "project_id": str(project.id),
        "model_id": str(model.id),
        "outline": project.outline,
    }


@router.get("/{project_id}/chapters", response_model=list[ChapterResponse])
def list_chapters(writer: WriterUserDep, db: DbSessionDep, project_id: UUID) -> list[Chapter]:
    project = _writer_project(db, writer, project_id)
    stmt = select(Chapter).where(Chapter.project_id == project.id).order_by(Chapter.order_index)
    return list(db.scalars(stmt).all())


@router.patch("/{project_id}/chapters/{chapter_id}", response_model=ChapterResponse)
def update_chapter(
    writer: WriterUserDep,
    db: DbSessionDep,
    project_id: UUID,
    chapter_id: UUID,
    payload: ChapterUpdate,
) -> Chapter:
    project = _writer_project(db, writer, project_id)
    chapter = db.get(Chapter, chapter_id)
    if chapter is None or chapter.project_id != project.id:
        raise HTTPException(status_code=404, detail="Chapter not found")
    updates = payload.model_dump(exclude_unset=True)
    for key, val in updates.items():
        setattr(chapter, key, val)
    if "content" in updates and chapter.content is not None:
        chapter.word_count = len(chapter.content)
        chapter.version += 1
    db.add(chapter)
    db.commit()
    db.refresh(chapter)
    return chapter


@router.post("/{project_id}/chapters/{chapter_id}/generate", response_model=ChapterResponse)
async def generate_chapter(
    writer: WriterUserDep,
    db: DbSessionDep,
    llm: LLMServiceDep,
    embedding: EmbeddingServiceDep,
    project_id: UUID,
    chapter_id: UUID,
    payload: ChapterGenerateRequest,
) -> Chapter:
    project = _writer_project(db, writer, project_id)
    chapter = db.get(Chapter, chapter_id)
    if chapter is None or chapter.project_id != project.id:
        raise HTTPException(status_code=404, detail="Chapter not found")

    model = _enabled_model_for_scenario(db, payload.model_id, project, "chapter_write")
    vectors = await embedding.embed_many([f"{project.title or project.topic or project.discipline}\n{chapter.title}"])
    chunks = VectorStore(db).search(project.id, vectors[0], limit=8) if vectors else []
    context = "\n\n".join([f"[参考片段 {idx + 1}]\n{c.content[:1200]}" for idx, c in enumerate(chunks)])

    content = await llm.call(
        user_id=writer.id,
        project_id=project.id,
        agent_name="chapter_writer",
        scenario="chapter_write",
        model=model,
        messages=[
            {
                "role": "system",
                "content": (
                    "你是中文学术论文代写工作台中的章节写作 Agent。"
                    "基于项目主题、章节标题和 RAG 参考片段生成严谨正文，避免编造不存在的引用。"
                ),
            },
            {
                "role": "user",
                "content": (
                    f"项目层次：{project.degree_level}\n"
                    f"学科：{project.discipline}\n"
                    f"论文题目：{project.title or '未定'}\n"
                    f"论文主题：{project.topic or '未定'}\n"
                    f"章节：{chapter.title}\n"
                    f"目标字数：{payload.target_words}\n\n"
                    f"学校/格式模板：\n{_school_context(db, project)}\n\n"
                    f"RAG 参考片段：\n{context or '暂无已确认片段，请先按通用学术结构写作。'}"
                ),
            },
        ],
    )
    if not content.strip():
        content = (
            f"{chapter.title}\n\n"
            f"本节围绕“{project.title or project.topic or project.discipline}”展开。"
            "当前为开发模式占位正文；接入 AI 中转站后将根据已确认 RAG 片段生成完整章节。"
        )
    chapter.content = content
    chapter.word_count = len(content)
    chapter.status = "generated"
    chapter.version += 1
    db.add(chapter)
    db.commit()
    db.refresh(chapter)
    return chapter


@router.post("/{project_id}/chapters/{chapter_id}/review", response_model=ChapterResponse)
async def review_chapter(
    writer: WriterUserDep,
    db: DbSessionDep,
    llm: LLMServiceDep,
    project_id: UUID,
    chapter_id: UUID,
    payload: ChapterReviewRequest,
) -> Chapter:
    project = _writer_project(db, writer, project_id)
    chapter = db.get(Chapter, chapter_id)
    if chapter is None or chapter.project_id != project.id:
        raise HTTPException(status_code=404, detail="Chapter not found")
    if not chapter.content:
        raise HTTPException(status_code=400, detail="Chapter has no content")

    model = _enabled_model_for_scenario(db, payload.model_id, project, "chapter_review")
    feedback = await llm.call(
        user_id=writer.id,
        project_id=project.id,
        agent_name="chapter_reviewer",
        scenario="chapter_review",
        model=model,
        messages=[
            {
                "role": "system",
                "content": (
                    "你是中文学术论文审校 Agent。请从结构、论证、学术表达、引用风险、"
                    "重复表达风险五方面给出精炼修改意见。"
                ),
            },
            {"role": "user", "content": f"章节：{chapter.title}\n\n正文：\n{chapter.content[:24000]}"},
        ],
    )
    chapter.feedback = feedback or "开发模式占位反馈：建议补充论据、压缩重复表述并完善章节衔接。"
    chapter.status = "reviewed"
    db.add(chapter)
    db.commit()
    db.refresh(chapter)
    return chapter


@router.post("/{project_id}/chapters/{chapter_id}/rewrite", response_model=ChapterResponse)
async def rewrite_chapter(
    writer: WriterUserDep,
    db: DbSessionDep,
    llm: LLMServiceDep,
    project_id: UUID,
    chapter_id: UUID,
    payload: ChapterRewriteRequest,
) -> Chapter:
    project = _writer_project(db, writer, project_id)
    chapter = db.get(Chapter, chapter_id)
    if chapter is None or chapter.project_id != project.id:
        raise HTTPException(status_code=404, detail="Chapter not found")
    if not chapter.content:
        raise HTTPException(status_code=400, detail="Chapter has no content")

    model = _enabled_model_for_scenario(db, payload.model_id, project, "chapter_rewrite")
    rewritten = await llm.call(
        user_id=writer.id,
        project_id=project.id,
        agent_name="chapter_rewriter",
        scenario="chapter_rewrite",
        model=model,
        messages=[
            {
                "role": "system",
                "content": (
                    "你是中文学术论文降重与润色 Agent。请在保持事实、论点和学术语气的前提下，"
                    "降低重复表达，改善行文流畅度。只输出改写后的正文。"
                ),
            },
            {
                "role": "user",
                "content": (
                    f"改写要求：{payload.instruction or chapter.feedback or '降低重复表达并润色'}\n\n"
                    f"章节：{chapter.title}\n\n原文：\n{chapter.content[:24000]}"
                ),
            },
        ],
    )
    if rewritten.strip():
        chapter.content = rewritten
        chapter.word_count = len(rewritten)
        chapter.version += 1
    chapter.status = "rewritten"
    db.add(chapter)
    db.commit()
    db.refresh(chapter)
    return chapter


@router.get("/{project_id}/export")
def export_project(
    writer: WriterUserDep,
    db: DbSessionDep,
    project_id: UUID,
    format: str = Query(default="markdown", pattern="^(markdown|latex|docx)$"),
) -> Response:
    project = _writer_project(db, writer, project_id)
    chapters = list(
        db.scalars(select(Chapter).where(Chapter.project_id == project.id).order_by(Chapter.order_index))
    )
    title = project.title or project.topic or f"{project.discipline}论文"
    school_context = _school_context(db, project)
    references = _reference_lines(db, project.id)

    if format == "latex":
        escaped_school_context = (
            school_context.replace("\\", "\\textbackslash{}").replace("{", "\\{").replace("}", "\\}")
        )
        body = "\n\n".join(
            [
                "\\section{" + ch.title.replace("\\", "\\textbackslash{}").replace("{", "\\{").replace("}", "\\}") + "}\n"
                + (ch.content or "")
                for ch in chapters
            ]
        )
        refs = "\n".join([f"\\item {ref}" for ref in references]) or "\\item 暂无参考文献"
        text = (
            "\\documentclass[UTF8]{ctexart}\n"
            "\\usepackage{geometry}\n"
            "\\geometry{a4paper, margin=2.5cm}\n"
            "\\title{" + title.replace("\\", "\\textbackslash{}").replace("{", "\\{").replace("}", "\\}") + "}\n"
            "\\begin{document}\n\\maketitle\n"
            "\\tableofcontents\n\\newpage\n"
            "\\begin{abstract}\n"
            f"{project.abstract or project.topic or '摘要待补充。'}\n"
            "\\end{abstract}\n"
            "\\section*{格式模板说明}\n"
            f"{escaped_school_context}\n\n"
            f"{body}\n"
            "\\begin{thebibliography}{99}\n"
            f"{refs}\n"
            "\\end{thebibliography}\n"
            "\\end{document}\n"
        )
        return Response(content=text, media_type="application/x-tex")

    if format == "docx":
        try:
            from docx import Document
        except ImportError as err:  # pragma: no cover
            raise HTTPException(status_code=500, detail="python-docx is not installed") from err

        doc = Document()
        doc.add_heading(title, level=0)
        doc.add_paragraph(f"层次：{project.degree_level}")
        doc.add_paragraph(f"学科：{project.discipline}")
        doc.add_paragraph("目录：请在 Word 中插入或更新自动目录。")
        doc.add_page_break()
        doc.add_heading("摘要", level=1)
        doc.add_paragraph(project.abstract or project.topic or "摘要待补充。")
        doc.add_heading("格式模板说明", level=1)
        for line in school_context.splitlines():
            doc.add_paragraph(line)
        for ch in chapters:
            doc.add_heading(ch.title, level=1)
            doc.add_paragraph(ch.content or "")
        doc.add_heading("参考文献", level=1)
        for ref in references or ["暂无参考文献"]:
            doc.add_paragraph(ref, style="List Number")
        buf = BytesIO()
        doc.save(buf)
        return Response(
            content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": 'attachment; filename="thesis.docx"'},
        )

    lines = [
        f"# {title}",
        "",
        f"- 层次：{project.degree_level}",
        f"- 学科：{project.discipline}",
        "",
        "## 目录",
        "",
        "请在最终排版工具中生成自动目录。",
        "",
        "## 摘要",
        "",
        project.abstract or project.topic or "摘要待补充。",
        "",
    ]
    lines.extend(["## 格式模板说明", school_context, ""])
    for ch in chapters:
        lines.extend([f"## {ch.title}", ch.content or "", ""])
    lines.extend(["## 参考文献", ""])
    for idx, ref in enumerate(references, start=1):
        lines.append(f"{idx}. {ref}")
    if not references:
        lines.append("暂无参考文献")
    return Response(content="\n".join(lines), media_type="text/markdown; charset=utf-8")

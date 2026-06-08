"""
Format-aware DOCX generation engine.

Reads TemplateFormatRules DSL and applies python-docx styles to produce
a publication-ready Word document that honours the school's thesis guidelines.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.shared import Cm, Emu, Pt, RGBColor
from docx.text.paragraph import Paragraph


_FONT_SIZE_MAP: dict[str, float] = {
    "初号": 42, "小初号": 36, "一号": 26, "小一号": 24,
    "二号": 22, "小二号": 18, "三号": 16, "小三号": 15,
    "四号": 14, "小四号": 12, "五号": 10.5, "小五号": 9,
    "六号": 7.5, "小六号": 6.5, "七号": 5.5, "八号": 5,
}

_ALIGN_MAP: dict[str, int] = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justified": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _pt(rules: dict, path: str, fallback: float) -> Pt:
    val = _deep(rules, path, fallback)
    return Pt(float(val))


def _cm(rules: dict, path: str, fallback: float) -> Cm:
    val = _deep(rules, path, fallback)
    return Cm(float(val))


def _deep(d: dict, path: str, fallback: Any = None) -> Any:
    keys = path.split(".")
    cur: Any = d
    for k in keys:
        if not isinstance(cur, dict):
            return fallback
        cur = cur.get(k)
        if cur is None:
            return fallback
    return cur


def _resolve_font_size(font_cfg: dict) -> float:
    if "size_pt" in font_cfg:
        return float(font_cfg["size_pt"])
    name = font_cfg.get("size_name", "")
    return _FONT_SIZE_MAP.get(name, 12)


def _set_run_font(run, family: str, size_pt: float, bold: bool = False):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.name = family
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), family)
    rfonts.set(qn("w:hAnsi"), family)
    rfonts.set(qn("w:eastAsia"), family)


def _set_paragraph_format(para: Paragraph, spacing_cfg: dict, font_cfg: dict | None = None):
    pf = para.paragraph_format
    line = spacing_cfg.get("line", 1.5)
    pf.line_spacing = float(line)

    after = spacing_cfg.get("paragraph_after", 0)
    if after:
        pf.space_after = Pt(float(after))
    before = spacing_cfg.get("paragraph_before", 0)
    if before:
        pf.space_before = Pt(float(before))

    indent = spacing_cfg.get("first_line_indent")
    if indent and font_cfg:
        size_pt = _resolve_font_size(font_cfg)
        pf.first_line_indent = Pt(size_pt * float(indent))

    if font_cfg:
        align_str = font_cfg.get("align", "")
        if align_str in _ALIGN_MAP:
            pf.alignment = _ALIGN_MAP[align_str]


def _add_styled_paragraph(
    doc: Document,
    text: str,
    font_cfg: dict,
    spacing_cfg: dict,
    extra_before: float = 0,
    extra_after: float = 0,
) -> Paragraph:
    para = doc.add_paragraph()
    _set_paragraph_format(para, spacing_cfg, font_cfg)
    if extra_before:
        para.paragraph_format.space_before = Pt(extra_before)
    if extra_after:
        para.paragraph_format.space_after = Pt(extra_after)

    run = para.add_run(text)
    family = font_cfg.get("family", "宋体")
    size_pt = _resolve_font_size(font_cfg)
    bold = font_cfg.get("bold", False)
    _set_run_font(run, family, size_pt, bold)

    align_str = font_cfg.get("align", "")
    if align_str in _ALIGN_MAP:
        para.alignment = _ALIGN_MAP[align_str]

    return para


def _add_header_footer(section, rules: dict):
    header_cfg = _deep(rules, "fonts.header", {})
    footer_cfg = _deep(rules, "fonts.footer", {})

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = hp.add_run()
    _set_run_font(
        hr,
        header_cfg.get("family", "宋体"),
        _resolve_font_size(header_cfg) if header_cfg else 9,
        header_cfg.get("bold", False),
    )

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_number_field(fp)
    fr = fp.runs[0] if fp.runs else fp.add_run()
    _set_run_font(
        fr,
        footer_cfg.get("family", "宋体"),
        _resolve_font_size(footer_cfg) if footer_cfg else 9,
        footer_cfg.get("bold", False),
    )


def _add_page_number_field(paragraph: Paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    run._element.append(fld_char_begin)

    run2 = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    run2._element.append(instr)

    run3 = paragraph.add_run()
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run3._element.append(fld_char_end)


def _setup_page(doc: Document, rules: dict):
    section = doc.sections[0]

    page = _deep(rules, "page", {})
    margin = _deep(rules, "margin", {})

    if page.get("orientation") == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Cm(29.7), Cm(21)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Cm(21), Cm(29.7)

    section.top_margin = _cm(rules, "margin.top", 2.5)
    section.bottom_margin = _cm(rules, "margin.bottom", 2.0)
    section.left_margin = _cm(rules, "margin.left", 2.5)
    section.right_margin = _cm(rules, "margin.right", 2.0)

    _add_header_footer(section, rules)


def _build_cover_page(doc: Document, rules: dict, project_meta: dict):
    cover_cfg = _deep(rules, "cover", {})
    fields = cover_cfg.get("fields", ["题名", "作者姓名", "学科门类", "指导教师", "完成日期"])

    title_font = _deep(rules, "fonts.chapter_title", {"family": "黑体", "size_pt": 16, "bold": True})
    body_font = _deep(rules, "fonts.body", {"family": "宋体", "size_pt": 12})
    spacing = _deep(rules, "spacing", {})

    for _ in range(3):
        doc.add_paragraph()

    title_text = project_meta.get("title", "")
    if title_text:
        _add_styled_paragraph(doc, title_text, title_font, spacing)
    else:
        _add_styled_paragraph(doc, "（论文题目）", title_font, spacing)

    doc.add_paragraph()

    for field_name in fields:
        if field_name == "题名" or field_name == "副标题":
            continue
        value = project_meta.get(field_name, "")
        line = f"{field_name}：{value}" if value else f"{field_name}：___________"
        _add_styled_paragraph(doc, line, body_font, spacing)

    doc.add_page_break()


def _build_abstract(doc: Document, rules: dict, abstract_cn: str, abstract_en: str, keywords_cn: str, keywords_en: str):
    fonts = _deep(rules, "fonts", {})
    spacing = _deep(rules, "spacing", {})

    ab_title_font = fonts.get("abstract_title", {"family": "黑体", "size_pt": 16, "bold": True, "align": "center"})
    ab_body_font = fonts.get("abstract_body", {"family": "宋体", "size_pt": 12})
    ab_en_body_font = fonts.get("abstract_en_body", {"family": "Times New Roman", "size_pt": 12})
    kw_label_font = fonts.get("keywords_label", {"family": "宋体", "size_pt": 14, "bold": True})

    _add_styled_paragraph(doc, "摘要", ab_title_font, spacing)
    if abstract_cn:
        _add_styled_paragraph(doc, abstract_cn, ab_body_font, spacing)
    if keywords_cn:
        _add_styled_paragraph(doc, f"关键词：{keywords_cn}", kw_label_font, spacing)
    doc.add_page_break()

    _add_styled_paragraph(doc, "Abstract", ab_title_font, spacing)
    if abstract_en:
        _add_styled_paragraph(doc, abstract_en, ab_en_body_font, spacing)
    if keywords_en:
        kw_en_label = fonts.get("keywords_en_label", {"family": "Times New Roman", "size_pt": 14, "bold": True})
        _add_styled_paragraph(doc, f"Keywords: {keywords_en}", kw_en_label, spacing)
    doc.add_page_break()


def _heading_font_key(level: int) -> str:
    if level == 1:
        return "chapter_title"
    return f"section_l{level - 1}"


def _add_chapter_heading(doc: Document, title: str, level: int, rules: dict):
    fonts = _deep(rules, "fonts", {})
    spacing = _deep(rules, "spacing", {})

    font_key = _heading_font_key(level)
    font_cfg = fonts.get(font_key, fonts.get("chapter_title", {"family": "黑体", "size_pt": 16, "bold": True}))

    before_key = f"{'chapter' if level == 1 else f'section_l{level-1}'}_before_lines"
    before_pts = float(spacing.get(before_key, 0)) * 12

    para = _add_styled_paragraph(doc, title, font_cfg, spacing, extra_before=before_pts)
    return para


def _add_body_text(doc: Document, text: str, rules: dict):
    fonts = _deep(rules, "fonts", {})
    spacing = _deep(rules, "spacing", {})
    body_font = fonts.get("body", {"family": "宋体", "size_pt": 12})

    for paragraph_text in text.split("\n"):
        stripped = paragraph_text.strip()
        if not stripped:
            continue

        if stripped.startswith("## "):
            _add_chapter_heading(doc, stripped[3:].strip(), 1, rules)
        elif stripped.startswith("### "):
            _add_chapter_heading(doc, stripped[4:].strip(), 2, rules)
        elif stripped.startswith("#### "):
            _add_chapter_heading(doc, stripped[5:].strip(), 3, rules)
        elif stripped.startswith("##### "):
            _add_chapter_heading(doc, stripped[6:].strip(), 4, rules)
        elif stripped.startswith("|") and "|" in stripped[1:]:
            para = doc.add_paragraph(stripped)
            pf = para.paragraph_format
            pf.line_spacing = float(spacing.get("line", 1.5))
            run = para.runs[0] if para.runs else para.add_run()
            _set_run_font(run, body_font.get("family", "宋体"), _resolve_font_size(body_font), False)
        else:
            _add_styled_paragraph(doc, stripped, body_font, spacing)


def _build_references_section(doc: Document, references: list[str], rules: dict):
    fonts = _deep(rules, "fonts", {})
    spacing = _deep(rules, "spacing", {})
    chapter_font = fonts.get("chapter_title", {"family": "黑体", "size_pt": 16, "bold": True, "align": "center"})
    body_font = fonts.get("body", {"family": "宋体", "size_pt": 12})

    _add_styled_paragraph(doc, "参考文献", chapter_font, spacing)

    for idx, ref in enumerate(references, start=1):
        text = f"[{idx}] {ref}"
        _add_styled_paragraph(doc, text, body_font, spacing)


class DocxService:
    """Generates a format-compliant DOCX from chapters + DSL rules."""

    def __init__(self, rules_json: dict | None = None, structure_json: dict | None = None, citation_json: dict | None = None):
        self.rules = rules_json or {}
        self.structure = structure_json or {}
        self.citation = citation_json or {}

    def generate(
        self,
        *,
        title: str,
        degree_level: str,
        discipline: str,
        abstract_cn: str = "",
        abstract_en: str = "",
        keywords_cn: str = "",
        keywords_en: str = "",
        chapters: list[dict[str, str]],
        references: list[str] | None = None,
        project_meta: dict | None = None,
    ) -> bytes:
        doc = Document()
        _setup_page(doc, self.rules)

        meta = project_meta or {}
        meta.setdefault("title", title)
        meta.setdefault("学科门类", discipline)

        degree_labels = {"bachelor": "本科", "master": "硕士", "doctor": "博士"}
        meta.setdefault("学位层次", degree_labels.get(degree_level, degree_level))

        _build_cover_page(doc, self.rules, meta)
        _build_abstract(doc, self.rules, abstract_cn, abstract_en, keywords_cn, keywords_en)

        # TOC placeholder
        fonts = _deep(self.rules, "fonts", {})
        spacing = _deep(self.rules, "spacing", {})
        toc_font = fonts.get("toc_title", {"family": "黑体", "size_pt": 16, "bold": True, "align": "center"})
        _add_styled_paragraph(doc, "目  录", toc_font, spacing)
        toc_para = doc.add_paragraph()
        run = toc_para.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        run._element.append(fld_begin)
        run2 = toc_para.add_run()
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        toc_depth = _deep(self.rules, "numbering.toc_depth", 2)
        instr.text = f' TOC \\o "1-{toc_depth}" \\h \\z \\u '
        run2._element.append(instr)
        run3 = toc_para.add_run()
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run3._element.append(fld_end)
        doc.add_paragraph("（请在 Word 中右键此处 → 更新域 以生成目录）")
        doc.add_page_break()

        for ch in chapters:
            ch_title = ch.get("title", "")
            ch_content = ch.get("content", "")
            _add_chapter_heading(doc, ch_title, 1, self.rules)
            if ch_content:
                _add_body_text(doc, ch_content, self.rules)

        if references:
            doc.add_page_break()
            _build_references_section(doc, references, self.rules)

        buf = BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def generate_as_doc(
        self,
        *,
        title: str,
        degree_level: str,
        discipline: str,
        abstract_cn: str = "",
        abstract_en: str = "",
        keywords_cn: str = "",
        keywords_en: str = "",
        chapters: list[dict[str, str]],
        references: list[str] | None = None,
        project_meta: dict | None = None,
    ) -> bytes:
        """
        Generate document in legacy .doc format.
        
        This method:
        1. Generates .docx using python-docx
        2. Converts to .doc using FormatConverter
        3. Returns .doc bytes
        """
        # First generate .docx
        docx_bytes = self.generate(
            title=title,
            degree_level=degree_level,
            discipline=discipline,
            abstract_cn=abstract_cn,
            abstract_en=abstract_en,
            keywords_cn=keywords_cn,
            keywords_en=keywords_en,
            chapters=chapters,
            references=references,
            project_meta=project_meta,
        )
        
        # Convert to .doc
        import tempfile
        from pathlib import Path
        from app.services.format_converter import get_converter
        
        converter = get_converter()
        
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp_docx:
            tmp_docx.write(docx_bytes)
            tmp_docx_path = Path(tmp_docx.name)
        
        try:
            tmp_doc_path = tmp_docx_path.with_suffix('.doc')
            converter.convert(tmp_docx_path, 'doc', tmp_doc_path)
            return tmp_doc_path.read_bytes()
        finally:
            tmp_docx_path.unlink(missing_ok=True)
            if 'tmp_doc_path' in locals():
                tmp_doc_path.unlink(missing_ok=True)
